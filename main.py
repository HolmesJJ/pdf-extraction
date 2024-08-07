import os
import layoutparser as lp

from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage

DATA_DIR = 'data'
OUTPUT_DIR = 'output'
INPUT_FILE = os.path.join(DATA_DIR, 'Sample.pdf')
OUTPUT_FILE = os.path.join(OUTPUT_DIR, 'Sample.xlsx')


def run():
    pdf_layouts, pdf_images = lp.load_pdf(INPUT_FILE, load_images=True)
    model = lp.models.Detectron2LayoutModel('lp://HJDataset/faster_rcnn_R_50_FPN_3x/config')
    print(model)
    wb = Workbook()
    ws = wb.active
    row = 1
    for page_number, (page_layout, page_image) in enumerate(zip(pdf_layouts, pdf_images), 1):
        layout = model.detect(page_layout)
        print(layout)
        for element in layout:
            print(element.type)
            if element.type == 'Text':
                segment_image = page_image.crop(element.coordinates)
                text = lp.ocr_agent.detect(segment_image)
                ws[f'A{row}'] = text
                ws[f'B{row}'] = f'Page {page_number}'
            elif element.type == 'Figure':
                img_path = os.path.join(OUTPUT_DIR, f'image_p{page_number}_{element.id}.png')
                segment_image = page_image.crop(element.coordinates)
                segment_image.save(img_path)
                img = OpenpyxlImage(img_path)
                ws.add_image(img, f'A{row}')
            row += 1
    wb.save(OUTPUT_FILE)
    wb.close()


if __name__ == '__main__':
    print(list(lp.models.auto_layoutmodel.ALL_AVAILABLE_BACKENDS))
    run()
